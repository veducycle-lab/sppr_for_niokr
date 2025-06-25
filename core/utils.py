import os
import docx
import PyPDF2
import openpyxl
from rdflib import Graph

def process_file(filepath):
    """
    Возвращает список текстовых фрагментов из файлов PDF, DOCX, XLSX.
    Лимитирует размер фрагмента ≤ 3000 символов.
    """
    ext = os.path.splitext(filepath)[-1].lower()
    max_chunk_size = 3000
    if ext == '.pdf':
        text = extract_text_from_pdf(filepath)
    elif ext in ('.docx', '.doc'):
        text = extract_text_from_docx(filepath)
    elif ext in ('.xlsx', '.xls'):
        text = extract_text_from_xlsx(filepath)
    else:
        raise ValueError("Некорректный формат файла!")

    # Разбиваем на куски длиной ≤ 3000 символов
    return fragment_text(text, max_chunk_size)

def fragment_text(text, max_size):
    """
    Делит длинный текст на части, не превышающие max_size символов.
    """
    fragments = []
    start = 0
    while start < len(text):
        end = start + max_size
        fragments.append(text[start:end])
        start = end
    return fragments

def extract_text_from_pdf(filepath):
    with open(filepath, 'rb') as f:
        reader = PyPDF2.PdfReader(f)
        all_text = ''
        for page in reader.pages:
            all_text += page.extract_text() or ''
        return all_text

def extract_text_from_docx(filepath):
    doc = docx.Document(filepath)
    all_text = ''
    for para in doc.paragraphs:
        all_text += para.text + '\n'
    return all_text

def extract_text_from_xlsx(filepath):
    wb = openpyxl.load_workbook(filepath, data_only=True)
    all_text = ''
    for sheet in wb.worksheets:
        for row in sheet.iter_rows():
            row_text = ' '.join([str(cell.value) if cell.value else '' for cell in row])
            all_text += row_text + '\n'
    return all_text

def load_ontology_from_owl(file_path):
    g = Graph()
    g.parse(file_path, format="xml")  # или "turtle" если ваш OWL - Turtle
    triplets = [(str(s), str(p), str(o)) for s, p, o in g.triples((None, None, None))]
    return triplets